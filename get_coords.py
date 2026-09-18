from pptx import Presentation
import os

pptx_path = '/Users/s1ncer3ly/Desktop/Video Template hackathon 3.pptx'
if not os.path.exists(pptx_path):
    print(f"File not found: {pptx_path}")
    exit()

prs = Presentation(pptx_path)
slide = prs.slides[7] # Page 8 (0-indexed)

print(f"Slide 8 shapes:")
for i, shape in enumerate(slide.shapes):
    if shape.has_text_frame:
        print(f"{i}: {shape.text} | Left: {shape.left}, Top: {shape.top}, Width: {shape.width}, Height: {shape.height}")
    else:
        print(f"{i}: [No Text] | Left: {shape.left}, Top: {shape.top}, Width: {shape.width}, Height: {shape.height}")
